"""
Talaba Yordamchi Bot (Groq + Pexels, Railway uchun tayyor)
-----------------------------------------------------------------
/slayd     - til -> mavzu -> fan -> soni -> dizayn -> bajaruvchi -> qabul qiluvchi
             -> ko'rib chiqish -> .pptx
/mustaqil  - til -> mavzu -> fan -> soni -> bajaruvchi -> qabul qiluvchi
             -> ko'rib chiqish -> .docx (+ imkon bo'lsa .pdf)

Qo'shimcha imkoniyatlar:
- Kunlik so'rov limiti va har-foydalanuvchi navbat (bitta vaqtda bitta so'rov)
- Umumiy vaqt chegarasi (bot "osilib" qolmasligi uchun)
- Fan/Bajaruvchi/Qabul qiluvchi ma'lumotlarini SQLite'da eslab qolish
- Xatolarni SQLite'ga yozish + ixtiyoriy admin xabarnomasi (ADMIN_CHAT_ID)
- "Boshqacha variant" (qayta generatsiya) tugmasi
"""

import os
import re
import json
import time
import random
import shutil
import sqlite3
import asyncio
import tempfile
import traceback
import subprocess
from io import BytesIO
from datetime import date

import requests
from groq import Groq
from telegram import Update, InlineKeyboardButton, InlineKeyboardMarkup
from telegram.ext import (
    ApplicationBuilder,
    CommandHandler,
    CallbackQueryHandler,
    ContextTypes,
    ConversationHandler,
    MessageHandler,
    filters,
)

from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE

from docx import Document
from docx.shared import Pt as DocxPt
from docx.enum.text import WD_ALIGN_PARAGRAPH

TELEGRAM_TOKEN = os.environ.get("TELEGRAM_TOKEN")
GROQ_API_KEY = os.environ.get("GROQ_API_KEY")
PEXELS_API_KEY = os.environ.get("PEXELS_API_KEY")
ADMIN_CHAT_ID = os.environ.get("ADMIN_CHAT_ID")  # ixtiyoriy: xatolarni shu chatga yuboradi

# DIQQAT: Railway'da doimiy volume ulanmasa, bu fayl konteyner qayta ishga
# tushganda o'chib ketadi (profil/limit ma'lumotlari yo'qoladi). Doimiy
# saqlash uchun Railway'da Volume yarating va DB_PATH'ni shu volume ichidagi
# yo'lga (masalan /data/bot_data.db) o'rnating.
DB_PATH = os.environ.get("DB_PATH", "bot_data.db")
DAILY_LIMIT = int(os.environ.get("DAILY_LIMIT", "5"))
GEN_TIMEOUT = int(os.environ.get("GEN_TIMEOUT_SEC", "420"))

if not TELEGRAM_TOKEN or not GROQ_API_KEY:
    raise ValueError(
        "TELEGRAM_TOKEN yoki GROQ_API_KEY topilmadi! "
        "Environment variables to'g'ri sozlanganini tekshiring."
    )

client = Groq(api_key=GROQ_API_KEY)
MODEL = "openai/gpt-oss-120b"

MAX_COUNT = 12
SONI_TANLOVLARI = [4, 6, 8, 10, 12]

# Suhbat holatlari
TIL, MAVZU, FAN, SONI, TEMA, BAJARUVCHI, QABUL, TASDIQ = range(8)

# Bitta vaqtda faqat bitta so'rov generatsiya qilinishi uchun (xotirada)
FAOL_FOYDALANUVCHILAR = set()

THEMES = [
    {"bg": RGBColor(0x1A, 0x23, 0x3A), "title": RGBColor(0xFF, 0xFF, 0xFF), "text": RGBColor(0xE0, 0xE0, 0xE0), "accent": RGBColor(0x4F, 0xA8, 0xE0)},
    {"bg": RGBColor(0xFF, 0xFF, 0xFF), "title": RGBColor(0x1A, 0x23, 0x3A), "text": RGBColor(0x33, 0x33, 0x33), "accent": RGBColor(0xE0, 0x6A, 0x4F)},
    {"bg": RGBColor(0x2D, 0x2A, 0x4A), "title": RGBColor(0xFF, 0xD9, 0x66), "text": RGBColor(0xF0, 0xF0, 0xF0), "accent": RGBColor(0xFF, 0xD9, 0x66)},
    {"bg": RGBColor(0xF4, 0xF1, 0xEA), "title": RGBColor(0x3A, 0x5A, 0x40), "text": RGBColor(0x2C, 0x2C, 0x2C), "accent": RGBColor(0x3A, 0x5A, 0x40)},
]
THEME_NAMES = ["🌌 Tungi ko'k", "☀️ Oq-parlaq", "🟣 Binafsha tun", "🌿 Tabiiy yashil"]

LANG = {
    "uz": {
        "label": "🇺🇿 O'zbek", "ai_lang": "o'zbek",
        "fan": "Fani", "bajaruvchi": "Bajaruvchi", "qabul": "Qabul qiluvchi",
        "kirish": "Kirish", "asosiy": "Asosiy qism", "xulosa": "Xulosa", "adabiyotlar": "Foydalanilgan adabiyotlar",
        "mavzu_savol": "Mavzuni kiriting:",
        "mavzu_xato": "Mavzu juda qisqa yoki noto'g'ri. Iltimos, kamida 3 ta harfdan iborat aniq mavzu kiriting:",
        "fan_savol": "Fan nomini kiriting (masalan: Iqtisodiyot nazariyasi):",
        "bajaruvchi_savol": "Bajaruvchi (talaba) F.I.Sh. kiriting:",
        "qabul_savol": "Qabul qiluvchi (fan o'qituvchisi) F.I.Sh. kiriting:",
    },
    "ru": {
        "label": "🇷🇺 Русский", "ai_lang": "рус",
        "fan": "Предмет", "bajaruvchi": "Исполнитель", "qabul": "Принял",
        "kirish": "Введение", "asosiy": "Основная часть", "xulosa": "Заключение", "adabiyotlar": "Использованная литература",
        "mavzu_savol": "Введите тему:",
        "mavzu_xato": "Тема слишком короткая. Введите тему из минимум 3 букв:",
        "fan_savol": "Введите название предмета (например: Экономическая теория):",
        "bajaruvchi_savol": "Введите Ф.И.О. исполнителя (студента):",
        "qabul_savol": "Введите Ф.И.О. преподавателя, принимающего работу:",
    },
    "en": {
        "label": "🇬🇧 English", "ai_lang": "ingliz",
        "fan": "Subject", "bajaruvchi": "Prepared by", "qabul": "Supervisor",
        "kirish": "Introduction", "asosiy": "Main part", "xulosa": "Conclusion", "adabiyotlar": "References",
        "mavzu_savol": "Enter the topic:",
        "mavzu_xato": "The topic is too short. Please enter a topic with at least 3 letters:",
        "fan_savol": "Enter the subject name (e.g. Economics):",
        "bajaruvchi_savol": "Enter the student's full name:",
        "qabul_savol": "Enter the supervising teacher's full name:",
    },
}


# ---------------------------------------------------------------------------
# SQLite: profil, kunlik limit, xatolar jurnali
# ---------------------------------------------------------------------------

def db_init():
    conn = sqlite3.connect(DB_PATH)
    conn.execute("""CREATE TABLE IF NOT EXISTS profiles (
        user_id INTEGER PRIMARY KEY, fan TEXT, bajaruvchi TEXT, qabul TEXT
    )""")
    conn.execute("""CREATE TABLE IF NOT EXISTS usage (
        user_id INTEGER, kun TEXT, soni INTEGER, PRIMARY KEY (user_id, kun)
    )""")
    conn.execute("""CREATE TABLE IF NOT EXISTS xatolar (
        id INTEGER PRIMARY KEY AUTOINCREMENT, vaqt TEXT, user_id INTEGER, joy TEXT, xato TEXT
    )""")
    conn.commit()
    conn.close()


def profil_olish(user_id: int):
    conn = sqlite3.connect(DB_PATH)
    row = conn.execute("SELECT fan, bajaruvchi, qabul FROM profiles WHERE user_id=?", (user_id,)).fetchone()
    conn.close()
    if row:
        return {"fan": row[0], "bajaruvchi": row[1], "qabul": row[2]}
    return None


def profil_saqlash(user_id: int, fan: str, bajaruvchi: str, qabul: str):
    conn = sqlite3.connect(DB_PATH)
    conn.execute(
        "INSERT INTO profiles (user_id, fan, bajaruvchi, qabul) VALUES (?, ?, ?, ?) "
        "ON CONFLICT(user_id) DO UPDATE SET fan=excluded.fan, bajaruvchi=excluded.bajaruvchi, qabul=excluded.qabul",
        (user_id, fan, bajaruvchi, qabul),
    )
    conn.commit()
    conn.close()


def bugungi_soni(user_id: int) -> int:
    kun = date.today().isoformat()
    conn = sqlite3.connect(DB_PATH)
    row = conn.execute("SELECT soni FROM usage WHERE user_id=? AND kun=?", (user_id, kun)).fetchone()
    conn.close()
    return row[0] if row else 0


def usage_oshirish(user_id: int):
    kun = date.today().isoformat()
    conn = sqlite3.connect(DB_PATH)
    conn.execute(
        "INSERT INTO usage (user_id, kun, soni) VALUES (?, ?, 1) "
        "ON CONFLICT(user_id, kun) DO UPDATE SET soni = soni + 1",
        (user_id, kun),
    )
    conn.commit()
    conn.close()


def xato_yozish(user_id: int, joy: str, xato):
    try:
        conn = sqlite3.connect(DB_PATH)
        conn.execute(
            "INSERT INTO xatolar (vaqt, user_id, joy, xato) VALUES (datetime('now'), ?, ?, ?)",
            (user_id, joy, str(xato)[:500]),
        )
        conn.commit()
        conn.close()
    except Exception:
        pass


async def adminga_xabar(context: ContextTypes.DEFAULT_TYPE, matn: str):
    if not ADMIN_CHAT_ID:
        return
    try:
        await context.bot.send_message(chat_id=ADMIN_CHAT_ID, text=matn[:4000])
    except Exception:
        pass


# ---------------------------------------------------------------------------
# Yordamchi funksiyalar
# ---------------------------------------------------------------------------

def ask_ai(prompt: str, max_tokens: int = 3000, max_retries: int = 3) -> str:
    """Groq'ga so'rov yuboradi. Vaqtinchalik rate-limit/overload xatolarida
    kichik kutish bilan qayta urinadi."""
    delay = 2
    last_err = None
    for attempt in range(max_retries):
        try:
            response = client.chat.completions.create(
                model=MODEL, max_tokens=max_tokens,
                messages=[{"role": "user", "content": prompt}],
            )
            return response.choices[0].message.content
        except Exception as e:
            last_err = e
            msg = str(e).lower()
            transient = any(x in msg for x in ["rate", "429", "overloaded", "503", "timeout"])
            if transient and attempt < max_retries - 1:
                time.sleep(delay)
                delay *= 2
                continue
            raise
    raise last_err


def strip_markdown(text: str) -> str:
    text = re.sub(r"\*\*(.*?)\*\*", r"\1", text)
    text = re.sub(r"\*(.*?)\*", r"\1", text)
    text = re.sub(r"^#{1,6}\s*", "", text, flags=re.MULTILINE)
    text = re.sub(r"^-\s+", "", text, flags=re.MULTILINE)
    return text


def extract_json(text: str) -> dict:
    text = re.sub(r"```(?:json)?", "", text)
    match = re.search(r"\{.*\}", text, re.DOTALL)
    if not match:
        raise ValueError("AI javobida JSON topilmadi")
    return json.loads(match.group(0))


def sanitize_filename(name: str) -> str:
    name = re.sub(r'[\\/:*?"<>|]', "", name).strip()
    return name[:60] if name else "hujjat"


def mavzu_yaroqlimi(matn: str) -> bool:
    matn = matn.strip()
    if len(matn) < 3:
        return False
    harflar = re.findall(r"[^\W\d_]", matn, flags=re.UNICODE)
    return len(harflar) >= 2


def get_image(query: str, per_page: int = 5):
    """Pexels'dan bir nechta natija olib, eng yuqori sifatlisini tanlaydi.
    Qaytaradi: (BytesIO yoki None, sabab_matni)"""
    if not PEXELS_API_KEY:
        return None, "PEXELS_API_KEY sozlanmagan (Railway Variables'da yo'q)"
    try:
        resp = requests.get(
            "https://api.pexels.com/v1/search",
            headers={"Authorization": PEXELS_API_KEY.strip()},
            params={"query": query, "per_page": per_page, "orientation": "landscape"},
            timeout=15,
        )
        if resp.status_code == 401:
            return None, "PEXELS_API_KEY noto'g'ri (401 - ruxsat berilmadi)"
        if resp.status_code != 200:
            return None, f"Pexels xatosi (status={resp.status_code})"
        photos = resp.json().get("photos", [])
        if not photos:
            return None, f"'{query}' uchun rasm topilmadi"
        best = max(photos, key=lambda ph: ph.get("width", 0))
        img_resp = requests.get(best["src"]["large"], timeout=15)
        if img_resp.status_code != 200:
            return None, "Rasm faylini yuklab bo'lmadi"
        return BytesIO(img_resp.content), None
    except Exception as e:
        return None, f"Kutilmagan xato: {e}"


def docx_to_pdf(docx_path: str, out_dir: str):
    """LibreOffice orqali docx'ni pdf'ga o'giradi. soffice topilmasa (odatiy
    Railway image'ida yo'q), None qaytaradi - bu holatda faqat docx yuboriladi."""
    soffice = shutil.which("soffice") or shutil.which("libreoffice")
    if not soffice:
        return None, "LibreOffice (soffice) o'rnatilmagan"
    try:
        result = subprocess.run(
            [soffice, "--headless", "--convert-to", "pdf", "--outdir", out_dir, docx_path],
            timeout=90, capture_output=True, text=True,
        )
        pdf_path = os.path.splitext(docx_path)[0] + ".pdf"
        if os.path.exists(pdf_path):
            return pdf_path, None
        return None, f"Konvertatsiya muvaffaqiyatsiz: {result.stderr[:300]}"
    except Exception as e:
        return None, f"PDF konvertatsiya xatosi: {e}"


def add_background(slide, prs, color):
    bg = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width, prs.slide_height)
    bg.fill.solid()
    bg.fill.fore_color.rgb = color
    bg.line.fill.background()
    bg.shadow.inherit = False
    return bg


def create_pptx(mavzu: str, fan: str, bajaruvchi: str, qabul: str, slides_data: list, til: str, tema_idx=None):
    L = LANG[til]
    prs = Presentation()
    prs.slide_width = Inches(13.33)
    prs.slide_height = Inches(7.5)
    theme = THEMES[tema_idx] if isinstance(tema_idx, int) and 0 <= tema_idx < len(THEMES) else random.choice(THEMES)
    blank = prs.slide_layouts[6]
    image_errors = []

    slide = prs.slides.add_slide(blank)
    add_background(slide, prs, theme["bg"])

    title_box = slide.shapes.add_textbox(Inches(1), Inches(2.0), Inches(11.33), Inches(1.5))
    tf = title_box.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = mavzu
    p.alignment = PP_ALIGN.CENTER
    p.font.size = Pt(40)
    p.font.bold = True
    p.font.color.rgb = theme["title"]

    info_lines = [f"{L['fan']}: {fan}", f"{L['bajaruvchi']}: {bajaruvchi}", f"{L['qabul']}: {qabul}"]
    info_box = slide.shapes.add_textbox(Inches(1), Inches(3.7), Inches(11.33), Inches(2.2))
    tf2 = info_box.text_frame
    tf2.word_wrap = True
    for i, line in enumerate(info_lines):
        p2 = tf2.paragraphs[0] if i == 0 else tf2.add_paragraph()
        p2.text = line
        p2.alignment = PP_ALIGN.CENTER
        p2.font.size = Pt(20)
        p2.font.color.rgb = theme["accent"]
        p2.space_after = Pt(6)

    total = len(slides_data)

    for idx, item in enumerate(slides_data, start=1):
        slide = prs.slides.add_slide(blank)
        add_background(slide, prs, theme["bg"])

        image_on_left = idx % 2 == 1
        text_x = Inches(6.8) if image_on_left else Inches(0.7)
        img_x = Inches(0.5) if image_on_left else Inches(7.1)

        title_box = slide.shapes.add_textbox(text_x, Inches(0.5), Inches(6.0), Inches(1.0))
        tf = title_box.text_frame
        tf.word_wrap = True
        p = tf.paragraphs[0]
        p.text = strip_markdown(item.get("title", ""))
        p.font.size = Pt(28)
        p.font.bold = True
        p.font.color.rgb = theme["title"]

        body_box = slide.shapes.add_textbox(text_x, Inches(1.6), Inches(6.0), Inches(5.3))
        tf = body_box.text_frame
        tf.word_wrap = True
        for i, bullet in enumerate(item.get("bullets", [])):
            p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
            p.text = f"•  {strip_markdown(bullet)}"
            p.font.size = Pt(13)
            p.font.color.rgb = theme["text"]
            p.space_after = Pt(8)

        query = item.get("image_query") or item.get("title") or mavzu
        img_data, err = get_image(query)
        if img_data:
            try:
                slide.shapes.add_picture(img_data, img_x, Inches(1.6), width=Inches(5.7), height=Inches(4.8))
            except Exception as e:
                image_errors.append(f"Slayd {idx}: rasmni joylashda xato - {e}")
        elif err:
            image_errors.append(f"Slayd {idx} ('{query}'): {err}")

        num_box = slide.shapes.add_textbox(Inches(12.5), Inches(7.05), Inches(0.7), Inches(0.4))
        p = num_box.text_frame.paragraphs[0]
        p.text = f"{idx}/{total}"
        p.font.size = Pt(12)
        p.font.color.rgb = theme["accent"]

    buffer = BytesIO()
    prs.save(buffer)
    buffer.seek(0)
    return buffer, image_errors


def create_docx(mavzu: str, fan: str, bajaruvchi: str, qabul: str, matn: str) -> BytesIO:
    doc = Document()
    for _ in range(5):
        doc.add_paragraph()
    title_p = doc.add_paragraph()
    title_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = title_p.add_run(mavzu)
    run.font.size = DocxPt(24)
    run.font.bold = True

    doc.add_paragraph()
    for label, value in [("Fan", fan), ("Bajaruvchi", bajaruvchi), ("Qabul qiluvchi", qabul)]:
        info_p = doc.add_paragraph()
        info_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = info_p.add_run(f"{label}: {value}")
        r.font.size = DocxPt(14)

    doc.add_page_break()

    matn = strip_markdown(matn)
    for paragraph in matn.split("\n"):
        paragraph = paragraph.strip()
        if not paragraph:
            continue
        if paragraph.endswith(":") and len(paragraph) < 40:
            doc.add_heading(paragraph, level=2)
        else:
            p = doc.add_paragraph(paragraph)
            for run in p.runs:
                run.font.size = DocxPt(12)

    buffer = BytesIO()
    doc.save(buffer)
    buffer.seek(0)
    return buffer


# ---------------------------------------------------------------------------
# AI generatsiya (bloklovchi funksiyalar - hammasi to_thread orqali chaqiriladi)
# ---------------------------------------------------------------------------

def generate_outline(mavzu: str, soni: int, til: str) -> list:
    ai_lang = LANG[til]["ai_lang"]
    prompt = (
        f"'{mavzu}' mavzusida {soni} ta slaydlik taqdimot uchun FAQAT slayd "
        f"sarlavhalari ro'yxatini tuz. FAQAT shu JSON formatida javob ber:\n"
        f'{{"titles": ["Sarlavha 1", "Sarlavha 2"]}}\n'
        f"Aynan {soni} ta sarlavha bo'lsin, mavzuni mantiqiy ketma-ketlikda yoritsin. "
        f"{ai_lang} tilida yoz."
    )
    data = extract_json(ask_ai(prompt, max_tokens=1000))
    titles = [t for t in data["titles"][:soni] if t and str(t).strip()]
    while len(titles) < soni:
        titles.append(f"{mavzu} - qo'shimcha ma'lumot")
    return titles


def generate_slide_content(mavzu: str, title: str, til: str) -> dict:
    ai_lang = LANG[til]["ai_lang"]
    prompt = (
        f"Mavzu: '{mavzu}'. Slayd sarlavhasi: '{title}'.\n"
        f"Shu slayd uchun FAQAT quyidagi JSON formatida javob ber:\n"
        f'{{"image_query": "2-3 word english keyword", "bullets": '
        f'["batafsil band 1", "batafsil band 2"]}}\n'
        f"4-5 ta band yoz, har biri kamida 20-25 so'zdan iborat bo'lsin. Bandlar {ai_lang} "
        f"tilida, image_query albatta inglizcha bo'lsin. Markdown belgilaridan foydalanma."
    )
    data = extract_json(ask_ai(prompt, max_tokens=1200))
    bullets = [b.strip() for b in data.get("bullets", []) if b and b.strip()]
    if len(bullets) < 2:
        raise ValueError(f"'{title}' uchun AI bo'sh/qisqa javob qaytardi")
    return {"image_query": (data.get("image_query") or "").strip() or mavzu, "bullets": bullets}


def generate_conclusion_slide(mavzu: str, til: str) -> dict:
    ai_lang = LANG[til]["ai_lang"]
    prompt = (
        f"'{mavzu}' mavzusidagi taqdimotning yakuniy xulosa slaydi uchun FAQAT shu JSON "
        f'formatida javob ber: {{"image_query": "2-3 word english keyword", "bullets": '
        f'["xulosa band 1", "xulosa band 2"]}}\n'
        f"3-4 ta band yoz, asosiy xulosalarni jamlab bersin. {ai_lang} tilida yoz."
    )
    data = extract_json(ask_ai(prompt, max_tokens=800))
    bullets = [b.strip() for b in data.get("bullets", []) if b and b.strip()]
    if not bullets:
        raise ValueError("Xulosa bo'sh qaytdi")
    return {"image_query": (data.get("image_query") or mavzu).strip(), "bullets": bullets}


def generate_references_slide(mavzu: str, til: str) -> dict:
    ai_lang = LANG[til]["ai_lang"]
    prompt = (
        f"'{mavzu}' mavzusi bo'yicha 4-5 ta namunaviy foydalanilgan adabiyotlar ro'yxatini "
        f'FAQAT shu JSON formatida ber: {{"bullets": ["1. Muallif F. Kitob nomi. Nashriyot, yil.", "2. ..."]}}\n'
        f"{ai_lang} tilida, akademik formatda yoz."
    )
    data = extract_json(ask_ai(prompt, max_tokens=600))
    bullets = [b.strip() for b in data.get("bullets", []) if b and b.strip()]
    if not bullets:
        raise ValueError("Adabiyotlar ro'yxati bo'sh qaytdi")
    return {"image_query": "books library", "bullets": bullets}


async def build_slides_data_async(mavzu: str, soni: int, til: str, status_msg, outline=None) -> list:
    L = LANG[til]
    titles = outline if outline else await asyncio.to_thread(generate_outline, mavzu, soni, til)
    slides_data = []

    for idx, title in enumerate(titles, start=1):
        content = None
        for _ in range(3):
            try:
                content = await asyncio.to_thread(generate_slide_content, mavzu, title, til)
                break
            except Exception:
                content = None
        if content is None:
            content = {
                "image_query": mavzu,
                "bullets": [
                    f"{title} - {mavzu} mavzusi doirasidagi muhim jihatlardan biri hisoblanadi.",
                    "Ushbu bo'lim bo'yicha batafsil ma'lumotni qo'shimcha manbalardan o'rganish tavsiya etiladi.",
                ],
            }
        slides_data.append({"title": title, "image_query": content.get("image_query", mavzu), "bullets": content.get("bullets", [])})
        try:
            await status_msg.edit_text(f"⏳ Slaydlar tayyorlanmoqda... {idx}/{soni}")
        except Exception:
            pass

    try:
        await status_msg.edit_text(f"⏳ {L['xulosa']} va {L['adabiyotlar'].lower()} tayyorlanmoqda...")
    except Exception:
        pass

    try:
        xulosa = await asyncio.to_thread(generate_conclusion_slide, mavzu, til)
    except Exception:
        xulosa = {"image_query": mavzu, "bullets": [f"{mavzu} mavzusi bo'yicha asosiy xulosalar shakllantirildi."]}
    slides_data.append({"title": L["xulosa"], "image_query": xulosa.get("image_query", mavzu), "bullets": xulosa.get("bullets", [])})

    try:
        adabiyotlar = await asyncio.to_thread(generate_references_slide, mavzu, til)
    except Exception:
        adabiyotlar = {"image_query": "books library", "bullets": ["Qo'shimcha manbalar mustaqil tarzda qo'shilishi tavsiya etiladi."]}
    slides_data.append({"title": L["adabiyotlar"], "image_query": adabiyotlar.get("image_query", "books"), "bullets": adabiyotlar.get("bullets", [])})

    return slides_data


# ---------------------------------------------------------------------------
# Yakuniy generatsiya va yuborish (asosiy oqim va "qayta generatsiya" ikkalasi
# ham shu funksiyalarni ishlatadi)
# ---------------------------------------------------------------------------

async def generate_slayd(context: ContextTypes.DEFAULT_TYPE, chat_id: int, user_id: int, params: dict):
    til = params["til"]
    L = LANG[til]
    status_msg = await context.bot.send_message(chat_id, f"⏳ Slaydlar tayyorlanmoqda... 0/{params['soni']}")
    try:
        async def ish():
            slides_data = await build_slides_data_async(params["mavzu"], params["soni"], til, status_msg, outline=params.get("outline"))
            try:
                await status_msg.edit_text("📦 Fayl yig'ilmoqda...")
            except Exception:
                pass
            return await asyncio.to_thread(
                create_pptx, params["mavzu"], params["fan"], params["bajaruvchi"], params["qabul"],
                slides_data, til, params.get("tema_idx"),
            )

        pptx_file, image_errors = await asyncio.wait_for(ish(), timeout=GEN_TIMEOUT)

        filename = sanitize_filename(params["mavzu"]) + ".pptx"
        pptx_file.name = filename
        await context.bot.send_document(chat_id, document=pptx_file, filename=filename)
        usage_oshirish(user_id)

        if image_errors:
            unique = list(dict.fromkeys(e.split(": ", 1)[-1] for e in image_errors))
            await context.bot.send_message(chat_id, "Eslatma: ba'zi slaydlarga rasm qo'shilmadi.\nSabab: " + "; ".join(unique[:3]))

        kb = InlineKeyboardMarkup([[InlineKeyboardButton("🔄 Boshqacha variant", callback_data="regen_slayd")]])
        await context.bot.send_message(
            chat_id,
            f"Eslatma: '{L['adabiyotlar']}' slaydidagi manbalar va matnlar sun'iy intellekt "
            f"yordamida tuzilgan - taqdimotdan foydalanishdan oldin ularni tekshirib, real "
            f"manbalar bilan almashtiring.",
            reply_markup=kb,
        )
    except asyncio.TimeoutError:
        await context.bot.send_message(chat_id, "⏱ Vaqt chegarasidan oshib ketdi, qaytadan urinib ko'ring.")
        xato_yozish(user_id, "slayd_timeout", "timeout")
    except Exception as e:
        traceback.print_exc()
        xato_yozish(user_id, "slayd", e)
        await adminga_xabar(context, f"Xato (slayd) user={user_id}: {e}")
        await context.bot.send_message(chat_id, f"Xatolik yuz berdi, qayta urinib ko'ring. ({e})")
    finally:
        FAOL_FOYDALANUVCHILAR.discard(user_id)


async def generate_mustaqil(context: ContextTypes.DEFAULT_TYPE, chat_id: int, user_id: int, params: dict):
    til = params["til"]
    L = LANG[til]
    status_msg = await context.bot.send_message(chat_id, "⏳ Mustaqil ish yozilmoqda, biroz kuting...")
    try:
        async def ish():
            soz_soni = params["soni"] * 350
            ai_lang = L["ai_lang"]
            prompt = (
                f"'{params['mavzu']}' mavzusida ({params['fan']} fanidan) chuqur va batafsil "
                f"mustaqil ish (referat) yoz, taxminan {soz_soni} so'z (bu {params['soni']} sahifaga teng). "
                f"{L['kirish']}:, {L['asosiy']}:, {L['xulosa']}:, {L['adabiyotlar']}: kabi bo'lim "
                f"sarlavhalari bilan. Asosiy qismni 2-3 ta kichik mavzuga bo'lib, har birini alohida "
                f"sarlavha bilan chuqur yorit. Har bir bo'limda aniq faktlar, misollar va tushuntirishlar "
                f"bo'lsin. {ai_lang} tilida, ilmiy uslubda, professional va ma'lumotga boy qilib yoz. "
                f"MUHIM: hech qanday Markdown belgilaridan (**, *, #, -) foydalanma."
            )
            matn = await asyncio.to_thread(ask_ai, prompt, 5000)
            try:
                await status_msg.edit_text("📦 Fayl yig'ilmoqda...")
            except Exception:
                pass
            return await asyncio.to_thread(create_docx, params["mavzu"], params["fan"], params["bajaruvchi"], params["qabul"], matn)

        docx_file = await asyncio.wait_for(ish(), timeout=GEN_TIMEOUT)
        filename = sanitize_filename(params["mavzu"]) + ".docx"

        with tempfile.TemporaryDirectory() as tmpdir:
            docx_path = os.path.join(tmpdir, filename)
            with open(docx_path, "wb") as f:
                f.write(docx_file.getvalue())
            with open(docx_path, "rb") as f:
                await context.bot.send_document(chat_id, document=f, filename=filename)
            usage_oshirish(user_id)

            pdf_path, pdf_err = await asyncio.to_thread(docx_to_pdf, docx_path, tmpdir)
            if pdf_path and os.path.exists(pdf_path):
                pdf_filename = sanitize_filename(params["mavzu"]) + ".pdf"
                with open(pdf_path, "rb") as f:
                    await context.bot.send_document(chat_id, document=f, filename=pdf_filename)
            elif pdf_err:
                xato_yozish(user_id, "pdf_convert", pdf_err)

        kb = InlineKeyboardMarkup([[InlineKeyboardButton("🔄 Boshqacha variant", callback_data="regen_mustaqil")]])
        await context.bot.send_message(
            chat_id,
            "Eslatma: matn sun'iy intellekt yordamida tuzilgan - taqdimdan oldin albatta o'qib "
            "chiqing va universitetingizning AI-matn siyosatini tekshiring.",
            reply_markup=kb,
        )
    except asyncio.TimeoutError:
        await context.bot.send_message(chat_id, "⏱ Vaqt chegarasidan oshib ketdi, qaytadan urinib ko'ring.")
        xato_yozish(user_id, "mustaqil_timeout", "timeout")
    except Exception as e:
        traceback.print_exc()
        xato_yozish(user_id, "mustaqil", e)
        await adminga_xabar(context, f"Xato (mustaqil) user={user_id}: {e}")
        await context.bot.send_message(chat_id, f"Xatolik yuz berdi, qayta urinib ko'ring. ({e})")
    finally:
        FAOL_FOYDALANUVCHILAR.discard(user_id)


# ---------------------------------------------------------------------------
# Telegram handlerlar
# ---------------------------------------------------------------------------

def til_klaviatura() -> InlineKeyboardMarkup:
    return InlineKeyboardMarkup([[InlineKeyboardButton(v["label"], callback_data=f"til_{k}")] for k, v in LANG.items()])


def soni_klaviatura() -> InlineKeyboardMarkup:
    row1 = [InlineKeyboardButton(str(n), callback_data=f"soni_{n}") for n in SONI_TANLOVLARI]
    return InlineKeyboardMarkup([row1, [InlineKeyboardButton("✍️ O'zim kiritaman", callback_data="soni_custom")]])


def tema_klaviatura() -> InlineKeyboardMarkup:
    return InlineKeyboardMarkup([[InlineKeyboardButton(name, callback_data=f"tema_{i}")] for i, name in enumerate(THEME_NAMES)])


def bajaruvchi_keyboard(user_id: int):
    p = profil_olish(user_id)
    if p and p.get("bajaruvchi"):
        return InlineKeyboardMarkup([[InlineKeyboardButton(f"✅ {p['bajaruvchi']}", callback_data="prof_bajaruvchi")]])
    return None


def qabul_keyboard(user_id: int):
    p = profil_olish(user_id)
    if p and p.get("qabul"):
        return InlineKeyboardMarkup([[InlineKeyboardButton(f"✅ {p['qabul']}", callback_data="prof_qabul")]])
    return None


FLOW_KALITLARI = ["turi", "til", "mavzu", "fan", "soni", "tema_idx", "bajaruvchi", "qabul", "_outline"]


def _flow_tozalash(context: ContextTypes.DEFAULT_TYPE):
    for k in FLOW_KALITLARI:
        context.user_data.pop(k, None)


async def start(update: Update, context: ContextTypes.DEFAULT_TYPE):
    await update.message.reply_text(
        "Salom! Men talabalarga yordam beruvchi botman.\n\n"
        "Buyruqlar:\n"
        "/slayd - rasmli, dizaynli PowerPoint (.pptx) tayyorlab beraman\n"
        "/mustaqil - Word (.docx) mustaqil ish tayyorlab beraman\n"
    )


async def slayd_start(update: Update, context: ContextTypes.DEFAULT_TYPE):
    user_id = update.effective_user.id
    if user_id in FAOL_FOYDALANUVCHILAR:
        await update.message.reply_text("Sizning oldingi so'rovingiz hali tugallanmadi, biroz kuting.")
        return ConversationHandler.END
    if bugungi_soni(user_id) >= DAILY_LIMIT:
        await update.message.reply_text(f"Bugungi limit ({DAILY_LIMIT} ta)ga yetdingiz. Ertaga qayta urinib ko'ring.")
        return ConversationHandler.END
    _flow_tozalash(context)
    context.user_data["turi"] = "slayd"
    await update.message.reply_text("Tilni tanlang:", reply_markup=til_klaviatura())
    return TIL


async def mustaqil_start(update: Update, context: ContextTypes.DEFAULT_TYPE):
    user_id = update.effective_user.id
    if user_id in FAOL_FOYDALANUVCHILAR:
        await update.message.reply_text("Sizning oldingi so'rovingiz hali tugallanmadi, biroz kuting.")
        return ConversationHandler.END
    if bugungi_soni(user_id) >= DAILY_LIMIT:
        await update.message.reply_text(f"Bugungi limit ({DAILY_LIMIT} ta)ga yetdingiz. Ertaga qayta urinib ko'ring.")
        return ConversationHandler.END
    _flow_tozalash(context)
    context.user_data["turi"] = "mustaqil"
    await update.message.reply_text("Tilni tanlang:", reply_markup=til_klaviatura())
    return TIL


async def handle_til(update: Update, context: ContextTypes.DEFAULT_TYPE):
    query = update.callback_query
    await query.answer()
    til = query.data.replace("til_", "")
    if til not in LANG:
        til = "uz"
    context.user_data["til"] = til
    await query.edit_message_text(LANG[til]["mavzu_savol"], reply_markup=None)
    return MAVZU


async def handle_mavzu(update: Update, context: ContextTypes.DEFAULT_TYPE):
    til = context.user_data.get("til", "uz")
    matn = update.message.text.strip()
    if not mavzu_yaroqlimi(matn):
        await update.message.reply_text(LANG[til]["mavzu_xato"])
        return MAVZU
    context.user_data["mavzu"] = matn
    p = profil_olish(update.effective_user.id)
    kb = InlineKeyboardMarkup([[InlineKeyboardButton(f"✅ {p['fan']}", callback_data="prof_fan")]]) if (p and p.get("fan")) else None
    await update.message.reply_text(LANG[til]["fan_savol"], reply_markup=kb)
    return FAN


async def handle_fan_callback(update: Update, context: ContextTypes.DEFAULT_TYPE):
    query = update.callback_query
    await query.answer()
    p = profil_olish(update.effective_user.id)
    context.user_data["fan"] = (p or {}).get("fan", "")
    await query.edit_message_text(f"Sahifalar soni (1 dan {MAX_COUNT} gacha):", reply_markup=soni_klaviatura())
    return SONI


async def handle_fan_text(update: Update, context: ContextTypes.DEFAULT_TYPE):
    context.user_data["fan"] = update.message.text.strip()
    await update.message.reply_text(f"Sahifalar soni (1 dan {MAX_COUNT} gacha):", reply_markup=soni_klaviatura())
    return SONI


async def _soni_keyingi_qadam(context, user_id, til):
    """SONI dan keyingi qadamni aniqlaydi: slayd -> TEMA, mustaqil -> BAJARUVCHI."""
    turi = context.user_data["turi"]
    if turi == "slayd":
        return "Taqdimot dizaynini tanlang:", tema_klaviatura(), TEMA
    return LANG[til]["bajaruvchi_savol"], bajaruvchi_keyboard(user_id), BAJARUVCHI


async def handle_soni_button(update: Update, context: ContextTypes.DEFAULT_TYPE):
    til = context.user_data.get("til", "uz")
    query = update.callback_query
    await query.answer()
    if query.data == "soni_custom":
        await query.edit_message_text(f"Nechta sahifa/slayd kerak, kiriting (1-{MAX_COUNT}):", reply_markup=None)
        return SONI
    context.user_data["soni"] = int(query.data.replace("soni_", ""))
    matn, kb, keyingi = await _soni_keyingi_qadam(context, update.effective_user.id, til)
    await query.edit_message_text(matn, reply_markup=kb)
    return keyingi


async def handle_soni_text(update: Update, context: ContextTypes.DEFAULT_TYPE):
    til = context.user_data.get("til", "uz")
    matn_raw = update.message.text.strip()
    if not matn_raw.isdigit() or not (1 <= int(matn_raw) <= MAX_COUNT):
        await update.message.reply_text(f"Iltimos, 1 dan {MAX_COUNT} gacha butun son kiriting:")
        return SONI
    context.user_data["soni"] = int(matn_raw)
    matn, kb, keyingi = await _soni_keyingi_qadam(context, update.effective_user.id, til)
    await update.message.reply_text(matn, reply_markup=kb)
    return keyingi


async def handle_tema_button(update: Update, context: ContextTypes.DEFAULT_TYPE):
    query = update.callback_query
    await query.answer()
    til = context.user_data.get("til", "uz")
    idx = int(query.data.replace("tema_", ""))
    context.user_data["tema_idx"] = idx if 0 <= idx < len(THEMES) else 0
    kb = bajaruvchi_keyboard(update.effective_user.id)
    await query.edit_message_text(LANG[til]["bajaruvchi_savol"], reply_markup=kb)
    return BAJARUVCHI


async def handle_bajaruvchi_callback(update: Update, context: ContextTypes.DEFAULT_TYPE):
    query = update.callback_query
    await query.answer()
    p = profil_olish(update.effective_user.id)
    context.user_data["bajaruvchi"] = (p or {}).get("bajaruvchi", "")
    til = context.user_data.get("til", "uz")
    kb = qabul_keyboard(update.effective_user.id)
    await query.edit_message_text(LANG[til]["qabul_savol"], reply_markup=kb)
    return QABUL


async def handle_bajaruvchi_text(update: Update, context: ContextTypes.DEFAULT_TYPE):
    context.user_data["bajaruvchi"] = update.message.text.strip()
    til = context.user_data.get("til", "uz")
    kb = qabul_keyboard(update.effective_user.id)
    await update.message.reply_text(LANG[til]["qabul_savol"], reply_markup=kb)
    return QABUL


async def tasdiqga_otish(context: ContextTypes.DEFAULT_TYPE, chat_id: int, user_data: dict, user_id: int):
    turi = user_data["turi"]
    til = user_data.get("til", "uz")
    L = LANG[til]
    matn = (
        "📋 Tekshiring:\n\n"
        f"Mavzu: {user_data['mavzu']}\n"
        f"{L['fan']}: {user_data['fan']}\n"
        f"Soni: {user_data['soni']}\n"
        f"{L['bajaruvchi']}: {user_data['bajaruvchi']}\n"
        f"{L['qabul']}: {user_data['qabul']}\n"
    )
    if turi == "slayd":
        try:
            titles = await asyncio.to_thread(generate_outline, user_data["mavzu"], user_data["soni"], til)
            user_data["_outline"] = titles
            matn += "\nSlayd sarlavhalari:\n" + "\n".join(f"{i+1}. {t}" for i, t in enumerate(titles))
        except Exception as e:
            xato_yozish(user_id, "outline_preview", e)
    kb = InlineKeyboardMarkup([[
        InlineKeyboardButton("✅ Tayyorlash", callback_data="tasdiq_ha"),
        InlineKeyboardButton("❌ Bekor qilish", callback_data="tasdiq_yoq"),
    ]])
    await context.bot.send_message(chat_id, matn, reply_markup=kb)


async def handle_qabul_callback(update: Update, context: ContextTypes.DEFAULT_TYPE):
    query = update.callback_query
    await query.answer()
    p = profil_olish(update.effective_user.id)
    context.user_data["qabul"] = (p or {}).get("qabul", "")
    await query.edit_message_text("Ma'lumotlar tayyorlanmoqda...", reply_markup=None)
    await tasdiqga_otish(context, update.effective_chat.id, context.user_data, update.effective_user.id)
    return TASDIQ


async def handle_qabul_text(update: Update, context: ContextTypes.DEFAULT_TYPE):
    context.user_data["qabul"] = update.message.text.strip()
    await update.message.reply_text("Ma'lumotlar tayyorlanmoqda...")
    await tasdiqga_otish(context, update.effective_chat.id, context.user_data, update.effective_user.id)
    return TASDIQ


async def handle_tasdiq(update: Update, context: ContextTypes.DEFAULT_TYPE):
    query = update.callback_query
    await query.answer()
    user_id = update.effective_user.id

    if query.data == "tasdiq_yoq":
        _flow_tozalash(context)
        await query.edit_message_text("Bekor qilindi.", reply_markup=None)
        return ConversationHandler.END

    if user_id in FAOL_FOYDALANUVCHILAR:
        await query.edit_message_text("Sizning boshqa so'rovingiz hali tugallanmadi, biroz kuting.", reply_markup=None)
        return ConversationHandler.END
    if bugungi_soni(user_id) >= DAILY_LIMIT:
        await query.edit_message_text(f"Kunlik limit ({DAILY_LIMIT} ta)ga yetdingiz. Ertaga qayta urinib ko'ring.", reply_markup=None)
        return ConversationHandler.END

    turi = context.user_data["turi"]
    params = {
        "mavzu": context.user_data["mavzu"],
        "fan": context.user_data["fan"],
        "soni": context.user_data["soni"],
        "bajaruvchi": context.user_data["bajaruvchi"],
        "qabul": context.user_data["qabul"],
        "til": context.user_data.get("til", "uz"),
        "tema_idx": context.user_data.get("tema_idx"),
        "outline": context.user_data.get("_outline"),
    }
    profil_saqlash(user_id, params["fan"], params["bajaruvchi"], params["qabul"])
    context.user_data[f"oxirgi_{turi}"] = params
    _flow_tozalash(context)

    FAOL_FOYDALANUVCHILAR.add(user_id)
    await query.edit_message_text("⏳ Boshlanmoqda...", reply_markup=None)
    chat_id = update.effective_chat.id
    if turi == "slayd":
        await generate_slayd(context, chat_id, user_id, params)
    else:
        await generate_mustaqil(context, chat_id, user_id, params)
    return ConversationHandler.END


async def handle_regenerate(update: Update, context: ContextTypes.DEFAULT_TYPE):
    query = update.callback_query
    await query.answer()
    turi = query.data.replace("regen_", "")
    params = context.user_data.get(f"oxirgi_{turi}")
    if not params:
        await query.message.reply_text("Avvalgi parametrlar topilmadi, iltimos /slayd yoki /mustaqil dan qaytadan boshlang.")
        return

    user_id = update.effective_user.id
    if user_id in FAOL_FOYDALANUVCHILAR:
        await query.message.reply_text("Sizning oldingi so'rovingiz hali tugallanmadi, biroz kuting.")
        return
    if bugungi_soni(user_id) >= DAILY_LIMIT:
        await query.message.reply_text(f"Kunlik limit ({DAILY_LIMIT} ta)ga yetdingiz. Ertaga qayta urinib ko'ring.")
        return

    FAOL_FOYDALANUVCHILAR.add(user_id)
    chat_id = update.effective_chat.id
    if turi == "slayd":
        await generate_slayd(context, chat_id, user_id, params)
    else:
        await generate_mustaqil(context, chat_id, user_id, params)


async def cancel(update: Update, context: ContextTypes.DEFAULT_TYPE):
    _flow_tozalash(context)
    await update.message.reply_text("Bekor qilindi.")
    return ConversationHandler.END


def main():
    db_init()
    app = ApplicationBuilder().token(TELEGRAM_TOKEN).build()

    conv = ConversationHandler(
        entry_points=[
            CommandHandler("slayd", slayd_start),
            CommandHandler("mustaqil", mustaqil_start),
        ],
        states={
            TIL: [CallbackQueryHandler(handle_til, pattern="^til_")],
            MAVZU: [MessageHandler(filters.TEXT & ~filters.COMMAND, handle_mavzu)],
            FAN: [
                CallbackQueryHandler(handle_fan_callback, pattern="^prof_fan$"),
                MessageHandler(filters.TEXT & ~filters.COMMAND, handle_fan_text),
            ],
            SONI: [
                CallbackQueryHandler(handle_soni_button, pattern="^soni_"),
                MessageHandler(filters.TEXT & ~filters.COMMAND, handle_soni_text),
            ],
            TEMA: [CallbackQueryHandler(handle_tema_button, pattern="^tema_")],
            BAJARUVCHI: [
                CallbackQueryHandler(handle_bajaruvchi_callback, pattern="^prof_bajaruvchi$"),
                MessageHandler(filters.TEXT & ~filters.COMMAND, handle_bajaruvchi_text),
            ],
            QABUL: [
                CallbackQueryHandler(handle_qabul_callback, pattern="^prof_qabul$"),
                MessageHandler(filters.TEXT & ~filters.COMMAND, handle_qabul_text),
            ],
            TASDIQ: [CallbackQueryHandler(handle_tasdiq, pattern="^tasdiq_")],
        },
        fallbacks=[
            CommandHandler("cancel", cancel),
            CommandHandler("slayd", slayd_start),
            CommandHandler("mustaqil", mustaqil_start),
        ],
        allow_reentry=True,
    )

    app.add_handler(CommandHandler("start", start))
    app.add_handler(conv)
    app.add_handler(CallbackQueryHandler(handle_regenerate, pattern="^regen_"))
    print("Bot ishga tushdi...")
    app.run_polling()


if __name__ == "__main__":
    main()
